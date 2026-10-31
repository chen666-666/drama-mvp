"""Freeze a truthful step-two delivery and stream checked media into a ZIP on disk."""
import io
import json
from pathlib import Path
import zipfile
from datetime import timezone

from models import ExportRun
from services.common import BEIJING, beijing, iso, utcnow
from services.downloads import DATA, catalog_detail, catalog_for, inside
from services.export import export_zip
from services.media import digest

FIELDS = ["ranking_id", "ranking_version", "records_cutoff", "position", "series_id", "series_title", "author_name",
          "video_id", "title", "original_url", "episode_number", "source_duration", "status", "access", "note", "evidence_url",
          "job_status", "attempts", "tool", "message", "relative_path", "size", "duration", "format", "sha256", "created_at",
          "validation_note", "manual_status", "manual_note", "checked_at", "file_in_package", "review_history"]


def word_report(ranking, catalogs, cutoff):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml.ns import qn
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(.75)
    section.left_margin = section.right_margin = Inches(.85)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = document.styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal = document.styles["Normal"]
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.3
    document.styles["Title"].font.size = Pt(23)
    document.styles["Heading 1"].font.size = Pt(16)
    document.styles["Heading 2"].font.size = Pt(12)
    document.add_heading("抖音 AI 短剧 Top3 下载交付报告", 0)
    document.add_paragraph("本报告记录固定榜单版本的入选作品、当前已发现公开正片及本地下载验收情况。名次沿用第一步结果，下载失败不改变名次。")
    document.add_paragraph(f"榜单 ID {ranking['id']}　批次 {ranking['run_id']}　版本 V{ranking['version']}\n下载记录截止 {beijing(cutoff)}")
    document.add_paragraph(ranking["disclaimer"])
    document.add_heading("榜单依据与范围", 1)
    document.add_paragraph(f"{ranking['title']}\n{ranking['explanation']}\n排名窗口截止 {ranking['cutoff_at_beijing']}\n榜单生成 {ranking['generated_at_beijing']}")
    if ranking.get("shortfall"):
        document.add_paragraph(ranking["shortfall"])
    for catalog in catalogs:
        item = next(row for row in ranking["items"] if row["selected"] and row["series_id"] == catalog["series_id"])
        summary = catalog["summary"]
        document.add_heading(f"第 {item['position']} 名 {item['series_title']}", 1)
        document.add_paragraph(f"原作者：{item['author_name']}　账号：{item['author_account_id']}\n作品 ID：{item['series_id']}　AI 类型：{item.get('ai_type') or '未记录'}\nAI 依据：{item.get('ai_basis') or '未记录'}")
        document.add_paragraph(f"代表集：{item['title']}\n{item['original_url']}\n采样：{item.get('sampled_at_beijing') or '未记录'}\n排序分：{item.get('score') if item.get('score') is not None else '保留官方名次'}")
        document.add_paragraph(f"已发现 {summary['discovered']} 集，确认公开 {summary['public']} 集，已下载 {summary['downloaded']} 集，自动校验通过 {summary['verified']} 集，人工抽检通过 {summary['manual_passed']} 集。")
        document.add_paragraph(f"{summary['completion']}；{summary['full_series']}。\n全剧总集数：{catalog['total_episodes'] if catalog['total_episodes'] is not None else '未知'}。\n目录说明：{catalog['note']}")
        gaps = [row for row in catalog["episodes"] if row["status"] != "approved" or row["access"] != "public" or not row.get("file_in_package") or not row["media"] or row["media"]["manual_status"] != "passed"]
        document.add_heading("逐集缺口与待办", 2)
        if not gaps:
            document.add_paragraph("本次已发现公开正片无下载或人工抽检缺口；全剧范围仍以目录核对说明为准。")
        for row in gaps:
            job, media = row["job"], row["media"]
            reason = row.get("package_error") or (job["message"] if job else "尚未下载")
            document.add_paragraph(f"视频 {row['video_id']}：{reason}；目录 {row['status']} / {row['access']}；人工抽检 {media['manual_status'] if media else '未开始'}。{row['note']}")
    if not catalogs:
        document.add_paragraph("该榜单版本没有合格入选作品，无视频交付。")
    document.add_heading("交付文件与验收口径", 1)
    document.add_paragraph("逐集来源、任务状态、文件大小、时长、SHA-256 与人工抽检记录见 04_剧集与下载记录.csv。来源与目录核验记录见 06_来源证据。视频位于 07_视频。")
    document.add_paragraph("仅将自动校验通过、文件仍存在且导出时摘要一致的视频收入资料包。自动校验与人工抽检分别记录；未完成抽检、缺集及总集数未知均不代表全剧完整。")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_package(session, ranking_row):
    from services.export import csv_text
    ranking = {"id": ranking_row.id, **json.loads(ranking_row.snapshot)}
    cutoff = utcnow()
    catalogs = [catalog_detail(session, catalog_for(session, ranking_row.id, row["series_id"])) for row in ranking["items"] if row["selected"]]
    rows, files = [], []
    for catalog in catalogs:
        item = next(row for row in ranking["items"] if row["selected"] and row["series_id"] == catalog["series_id"])
        for episode in catalog["episodes"]:
            job, media = episode["job"], episode["media"]
            episode["file_in_package"] = ""
            if media and job["status"] == "verified" and episode["status"] == "approved" and episode["access"] == "public" and media["manual_status"] != "rejected":
                path = inside(DATA, media["relative_path"])
                if path.is_file() and path.stat().st_size == media["size"] and digest(path) == media["sha256"]:
                    relative = Path(media["relative_path"]).parts[2:]
                    arcname = "07_视频/" + "/".join(relative)
                    episode["file_in_package"] = arcname
                    files.append((path, arcname))
                else:
                    episode["package_error"] = "文件缺失或摘要变化，未收入资料包"
                    media["exists"] = False
            rows.append({"ranking_id": ranking_row.id, "ranking_version": ranking["version"], "records_cutoff": iso(cutoff),
                         "position": item["position"], "series_id": item["series_id"], "series_title": item["series_title"], "author_name": item["author_name"],
                         **{key: episode.get(key) for key in FIELDS if key in episode},
                         "job_status": job["status"] if job else "not_queued", "attempts": job["attempts"] if job else 0,
                         "tool": job["tool"] if job else "", "message": episode.get("package_error") or (job["message"] if job else "尚未下载"),
                         **({key: media[key] for key in FIELDS if key in media} if media else {})})
        # Recount from the exact files delivered, including missing/tampered file exclusions.
        good = [row for row in catalog["episodes"] if row["file_in_package"]]
        catalog["summary"].update(verified=len(good), manual_passed=sum(row["media"]["manual_status"] == "passed" for row in good))
        if len(good) != catalog["summary"]["public"] or not good:
            catalog["summary"]["completion"] = "公开正片尚未下载完"
            catalog["summary"]["full_series"] = "全剧完整性未确认"
    export = ExportRun(ranking_id=ranking_row.id, cutoff_at=cutoff, relative_path="", summary="{}")
    session.add(export)
    session.flush()
    date = cutoff.replace(tzinfo=timezone.utc).astimezone(BEIJING).strftime("%Y-%m-%d")
    directory = f"抖音AI短剧_{date}"
    relative = f"exports/{directory}_榜单{ranking_row.id}_交付{export.id}.zip"
    path = inside(DATA, relative)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".tmp")
    try:
        with zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED, allowZip64=True) as archive:
            with zipfile.ZipFile(io.BytesIO(export_zip(ranking))) as first:
                for name in first.namelist():
                    archive.writestr(f"{directory}/{name}", first.read(name))
            archive.writestr(f"{directory}/01_榜单报告.docx", word_report(ranking, catalogs, cutoff))
            archive.writestr(f"{directory}/04_剧集与下载记录.csv", csv_text(rows, FIELDS))
            archive.writestr(f"{directory}/06_来源证据/第二步目录与核验.json", json.dumps({"ranking_id": ranking_row.id, "records_cutoff": iso(cutoff), "catalogs": catalogs}, ensure_ascii=False, indent=2))
            archive.writestr(f"{directory}/07_视频/", "")
            for file, name in files:
                archive.write(file, f"{directory}/{name}", compress_type=zipfile.ZIP_STORED)
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)
    export.relative_path, export.summary = relative, json.dumps({"files": len(files), "episodes": len(rows), "catalogs": catalogs}, ensure_ascii=False)
    return export
